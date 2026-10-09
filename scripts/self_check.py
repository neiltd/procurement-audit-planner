"""Step 5: independent re-derivation of the first flagged item of every test that ran.

Independent path: openpyxl for workbooks, csv module for contracts.csv, its own name normaliser,
thresholds read straight from rules.json. Writes the result into outputs/analytics.json.
Exit code non-zero on any mismatch.
"""
import csv
import json
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs"
FULL = OUT / "analytics-full"
SPEND = ROOT / "data" / "spend"
CONTRACTS = ROOT / "data" / "contracts.csv"
COLS = ["Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name"]

_wb_cache, _hdr_cache = {}, {}
results = []


def norm(name):
    s = re.sub(r"[^A-Za-z0-9\s]|_", " ", str(name or "").upper())
    return " ".join(w for w in s.split() if w not in {"LTD", "LIMITED", "PLC", "LLP"})


def sheet(fname):
    if fname not in _wb_cache:
        wb = openpyxl.load_workbook(SPEND / fname, read_only=False, data_only=True)
        ws = wb["Data to publish"]
        hdr = None
        for r in range(1, 51):
            vals = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
            if all(c in vals for c in COLS):
                hdr = {c: vals.index(c) + 1 for c in COLS}
                break
        assert hdr, fname
        _wb_cache[fname], _hdr_cache[fname] = ws, hdr
    return _wb_cache[fname], _hdr_cache[fname]


def raw_row(row_id):
    fname, sname, r = row_id.rsplit(":", 2)
    ws, hdr = sheet(fname)
    r = int(r)
    return {c: ws.cell(row=r, column=hdr[c]).value for c in COLS}


def clean_rows():
    with open(FULL / "spend_clean.csv", newline="") as f:
        return {r["row_id"]: r for r in csv.DictReader(f)}


def first_row(csvname):
    with open(FULL / csvname, newline="") as f:
        return next(csv.DictReader(f), None)


def confirm_raw(row_ids, clean):
    """Raw cell vs cleaned row: supplier, amount, date. Returns list of raw dicts and mismatches."""
    raws, bad = [], []
    for rid in row_ids:
        raw = raw_row(rid)
        c = clean[rid]
        if str(raw["Supplier name"]).strip() != c["Supplier name"]:
            bad.append(f"{rid} supplier raw={raw['Supplier name']!r} clean={c['Supplier name']!r}")
        if round(float(raw["Net amount"]), 2) != round(float(c["Net amount"]), 2):
            bad.append(f"{rid} amount raw={raw['Net amount']} clean={c['Net amount']}")
        rd = raw["Date"] if isinstance(raw["Date"], datetime) else datetime.strptime(str(raw["Date"]).strip(), "%m/%d/%y")
        cd = datetime.fromisoformat(c["Date"])
        if abs((rd - cd).total_seconds()) > 0.001:
            bad.append(f"{rid} date raw={rd} clean={cd}")
        if norm(raw["Supplier name"]) != c["supplier_norm"]:
            bad.append(f"{rid} norm self={norm(raw['Supplier name'])!r} clean={c['supplier_norm']!r}")
        raws.append({"rid": rid, "amount": float(raw["Net amount"]), "date": rd, "supplier": norm(raw["Supplier name"])})
    return raws, bad


def report(test, ok, detail):
    status = "OK" if ok else "MISMATCH"
    results.append((test, status, detail))
    print(f"{test}: {status} {detail}")


def contracts_rows():
    with open(CONTRACTS, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    return [r for r in rows if r["Organisation Name"].strip().upper() in {"WEST BERKSHIRE COUNCIL", "WEST BERKSHIRE"}]


def notice_suppliers(r):
    s = r["Supplier [Name|Address|Ref type|Ref Number|Is SME|Is VCSE]"].strip()
    if not s:
        return []
    return [norm(b.split("|")[0]) for b in s.strip("[]").split("][")]


def supplier_population_check(test, item, clean, factor, thr_value, thr_incl, compare="meets"):
    ids = item["row_ids"].split(";")
    raws, bad = confirm_raw(ids, clean)
    sups = {r["supplier"] for r in raws}
    if sups != {item["supplier_norm"]}:
        bad.append(f"suppliers in rows {sups} != {item['supplier_norm']}")
    total = round(sum(r["amount"] for r in raws), 2)
    if abs(total - float(item["total_3m"])) > 0.01:
        bad.append(f"total raw={total} list={item['total_3m']}")
    ann = total * factor
    if abs(ann - float(item["annualised_estimate"])) > 0.05:
        bad.append(f"annualised raw={ann:.2f} list={item['annualised_estimate']}")
    meets = ann >= thr_value if thr_incl else ann > thr_value
    if compare == "meets" and not meets:
        bad.append(f"annualised {ann:.2f} does not meet threshold {thr_value}")
    return raws, total, ann, bad


def main():
    a = json.loads((OUT / "analytics.json").read_text())
    rules = json.loads((OUT / "rules.json").read_text())
    clean = clean_rows()
    factor = 12 / a["window"]["months"]
    ws = datetime.fromisoformat(a["window"]["start"])
    by_clause = {(r.get("clause"), r.get("type")): r for r in rules["rules"]}
    q = by_clause[("App B row B1", "procedure")]
    p = by_clause[("App B row B1", "publication")]
    if a["thresholds"]["quote_threshold"]["value"] != q["threshold_low"] or \
       a["thresholds"]["publication_threshold"]["value"] != p["threshold_low"]:
        report("thresholds", False, "analytics thresholds differ from rules.json")
    tiers = {t["value"]: t["low_inclusive"] for t in a["thresholds"]["tiers"]}
    # tier values must each be cited in rules.json
    for v in tiers:
        cited = [r["rule_id"] for r in rules["rules"] if r.get("type") in {"procedure", "approval", "anti-avoidance"}
                 and v in (r.get("threshold_low"), r.get("threshold_high"))]
        if not cited:
            report("thresholds", False, f"tier {v} not in rules.json")

    tests = a["tests"]

    # 1 threshold_clustering
    t = tests["threshold_clustering"]
    if t["status"] == "ran":
        it = first_row("threshold_clustering.csv")
        raws, bad = confirm_raw([it["row_id"]], clean)
        T = float(it["tier"])
        incl = tiers[int(T)]
        amt = raws[0]["amount"]
        inband = (0.9 * T <= amt < T) if incl is True else (0.9 * T <= amt <= T)
        if not inband:
            bad.append(f"amount {amt} not in band for tier {T} (incl={incl})")
        report("threshold_clustering", not bad, f"{it['row_id']} {amt} in band of {T:.0f}" if not bad else "; ".join(bad))
    else:
        print("threshold_clustering: SKIPPED", t.get("reason"))

    # 2 split_purchases
    t = tests["split_purchases"]
    if t["status"] == "ran" and t["counts"]["total"]:
        it = first_row("split_purchases.csv")
        ids = it["row_ids"].split(";")
        raws, bad = confirm_raw(ids, clean)
        T = float(it["tier"])
        incl = tiers[int(T)]
        total = sum(r["amount"] for r in raws)
        days = [r["date"].date() for r in raws]
        if (max(days) - min(days)).days > 30:
            bad.append(f"span {(max(days) - min(days)).days} days > 30")
        for r in raws:
            if not (r["amount"] < T if incl is not False else r["amount"] <= T):
                bad.append(f"{r['rid']} {r['amount']} not below tier {T}")
        if not (total >= T if incl is not False else total > T):
            bad.append(f"group total {total:.2f} does not reach {T}")
        if abs(total - float(it["group_total"])) > 0.01:
            bad.append(f"total raw={total:.2f} list={it['group_total']}")
        if len({r["supplier"] for r in raws}) != 1 or len(raws) < 2:
            bad.append("group not a single supplier with >=2 payments")
        report("split_purchases", not bad,
               f"{it['supplier_norm']}: {len(raws)} payments total {total:.2f} >= {T:.0f} within {(max(days) - min(days)).days}d"
               if not bad else "; ".join(bad))
    else:
        print("split_purchases: SKIPPED", t.get("reason", "no groups"))

    # live notice names, independently
    live_names, all_names, crow = set(), set(), []
    if CONTRACTS.exists():
        crow = contracts_rows()
        for r in crow:
            end = r["End Date"].strip()
            is_live = (not end) or datetime.strptime(end, "%d/%m/%Y") >= ws
            for n in notice_suppliers(r):
                all_names.add(n)
                if is_live:
                    live_names.add(n)

    # 3 off_contract
    t = tests["off_contract"]
    if t["status"] == "ran":
        it = first_row("off_contract.csv")
        raws, total, ann, bad = supplier_population_check("off_contract", it, clean, factor, p["threshold_low"],
                                                         p.get("threshold_low_inclusive"))
        if it["supplier_norm"] in live_names:
            bad.append("supplier has an exact live notice match in contracts.csv")
        report("off_contract", not bad,
               f"{it['supplier_norm']}: annualised {ann:.2f} vs {p['threshold_low']} ({p['rule_id']}), no live notice"
               if not bad else "; ".join(bad))
    else:
        print("off_contract: SKIPPED", t.get("reason"))

    # 4 duplicates
    t = tests["duplicates"]
    if t["status"] == "ran":
        it = first_row("duplicates.csv")
        ids = it["row_ids"].split(";")
        raws, bad = confirm_raw(ids, clean)
        if len({r["supplier"] for r in raws}) != 1:
            bad.append("multiple suppliers")
        if len({round(r["amount"], 2) for r in raws}) != 1:
            bad.append("amounts differ")
        ds = sorted({r["date"] for r in raws})
        if it["type"] == "same_timestamp" and len(ds) != 1:
            bad.append("timestamps differ in same_timestamp item")
        if it["type"] == "different_dates":
            gaps = [(ds[i + 1].date() - ds[i].date()).days for i in range(len(ds) - 1)]
            if len(ds) < 2 or max(gaps) > 7:
                bad.append(f"gaps {gaps}")
        report("duplicates", not bad,
               f"{it['type']} {it['supplier_norm']}: {len(raws)} x {raws[0]['amount']:.2f}" if not bad else "; ".join(bad))
    else:
        print("duplicates: SKIPPED", t.get("reason"))

    # 5 expired_notice_spend
    t = tests["expired_notice_spend"]
    if t["status"] == "ran" and t["counts"]["total"]:
        it = first_row("expired_notice_spend.csv")
        ids = it["row_ids"].split(";")
        raws, bad = confirm_raw(ids, clean)
        total = sum(r["amount"] for r in raws)
        if abs(total - float(it["total_3m"])) > 0.01:
            bad.append(f"total raw={total:.2f} list={it['total_3m']}")
        ends = []
        for r in crow:
            if it["supplier_norm"] in notice_suppliers(r):
                e = r["End Date"].strip()
                ends.append((r["Notice Identifier"], datetime.strptime(e, "%d/%m/%Y") if e else None))
        if not ends:
            bad.append("no matching notice in contracts.csv")
        if any(e is None or e >= ws for _, e in ends):
            bad.append(f"a matched notice is live: {ends}")
        if {n for n, _ in ends} != set(it["notices"].split(";")):
            bad.append(f"notices raw={sorted(n for n, _ in ends)} list={it['notices']}")
        latest = max(e for _, e in ends) if ends and all(e for _, e in ends) else None
        report("expired_notice_spend", not bad,
               f"{it['supplier_norm']}: latest end {latest.date() if latest else None} < {ws.date()}, paid {total:.2f}"
               if not bad else "; ".join(bad))
    else:
        print("expired_notice_spend: SKIPPED", t.get("reason", "no items"))

    # 6 spend_vs_award
    t = tests["spend_vs_award"]
    if t["status"] == "ran" and t["counts"]["total"]:
        it = first_row("spend_vs_award.csv")
        ids = it["row_ids"].split(";")
        raws, bad = confirm_raw(ids, clean)
        total = sum(r["amount"] for r in raws)
        ann = total * factor
        nrow = [r for r in crow if r["Notice Identifier"] == it["notice_id"]]
        if len(nrow) != 1:
            bad.append(f"notice rows found: {len(nrow)}")
        else:
            r = nrow[0]
            sd = datetime.strptime(r["Start Date"].strip(), "%d/%m/%Y")
            ed = datetime.strptime(r["End Date"].strip(), "%d/%m/%Y")
            av = float(r["Awarded Value"].replace(",", ""))
            years = (ed - sd).days / 365.25
            ratio = ann / (av / years)
            sups = notice_suppliers(r)
            if sups != [it["supplier_norm"]]:
                bad.append(f"notice suppliers {sups}")
            if ed < ws:
                bad.append("notice not live")
            live_hits = [x for x in crow if it["supplier_norm"] in notice_suppliers(x)
                         and (not x["End Date"].strip() or datetime.strptime(x["End Date"].strip(), "%d/%m/%Y") >= ws)]
            if len({x["Notice Identifier"] for x in live_hits}) != 1:
                bad.append(f"supplier matches {len(live_hits)} live notices")
            if abs(av - float(it["awarded_value"])) > 0.01:
                bad.append(f"award raw={av} list={it['awarded_value']}")
            if abs(ratio - float(it["ratio"])) > 1e-6 * max(1, ratio):
                bad.append(f"ratio raw={ratio:.6f} list={it['ratio']}")
            if ratio < 1.25:
                bad.append(f"ratio {ratio:.3f} < 1.25")
            if not bad:
                report("spend_vs_award", True,
                       f"{it['supplier_norm']} {it['notice_id'][:12]}: end {ed.date()}, award {av:.2f}, ratio {ratio:.3f}")
        if bad:
            report("spend_vs_award", False, "; ".join(bad))
    else:
        print("spend_vs_award: SKIPPED", t.get("reason", "no items"))

    # 7 high_value_suppliers
    t = tests["high_value_suppliers"]
    it = first_row("high_value_suppliers.csv")
    raws, total, ann, bad = supplier_population_check("high_value_suppliers", it, clean, factor, q["threshold_low"],
                                                     q.get("threshold_low_inclusive"))
    if CONTRACTS.exists():
        exp = "exact" if it["supplier_norm"] in live_names else None
        if exp == "exact" and it["notice_match"] != "exact":
            bad.append(f"notice_match list={it['notice_match']} raw=exact")
        if exp is None and it["notice_match"] == "exact":
            bad.append("notice_match exact but no live notice found")
    report("high_value_suppliers", not bad,
           f"{it['supplier_norm']}: total {total:.2f}, annualised {ann:.2f} vs {q['threshold_low']} ({q['rule_id']}), match {it['notice_match']}"
           if not bad else "; ".join(bad))

    ok = all(s == "OK" for _, s, _ in results)
    print("PASS" if ok else "FAIL")
    a["self_check"] = {"script": "scripts/self_check.py", "result": "PASS" if ok else "FAIL",
                       "checked_at": datetime.now().isoformat(timespec="seconds"),
                       "lines": [f"{t_}: {s} {d}" for t_, s, d in results]}
    (OUT / "analytics.json").write_text(json.dumps(a, indent=2, default=str))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
